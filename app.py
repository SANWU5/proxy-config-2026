from __future__ import annotations

import csv
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - shown to the user at runtime
    Workbook = None
    load_workbook = None


if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent
CONFIG_PATH = APP_DIR / "config.json"

DEFAULT_SOURCE_DIR = Path(
    r"C:\Users\Administrator\AppData\Roaming\ziniaobrowserdatas\ziniao browser"
)
DEFAULT_OUTPUT_DIR = APP_DIR / "输出"
DEFAULT_OPERATOR = "未匹配运营"

INVALID_FILENAME_CHARS = r'<>:"/\|?*'


@dataclass(frozen=True)
class FileRule:
    pattern: str
    output_folder: str
    template: str


FILE_RULES = [
    FileRule("订单列表[", "订单明细", "{store}-订单明细-{operator}{suffix}"),
    FileRule("全店托管-订单报表", "广告明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("爆品打造-订单报表", "广告明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("商品推广-订单报表", "广告明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("关键词推广-订单报表", "广告明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("老用户推广-订单报表", "广告明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("账户报表", "广告消耗明细", "{store}-{stem}-{operator}{suffix}"),
    FileRule("refundOrders", "退款明细", "{store}-退款明细-{operator}{suffix}"),
]


def safe_filename(value: str) -> str:
    cleaned = "".join("_" if ch in INVALID_FILENAME_CHARS else ch for ch in value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    return cleaned or "未命名"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}({index}){path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"无法生成不重名文件：{path}")


def normalize_header(value: object) -> str:
    return str(value or "").strip().replace(" ", "").replace("\u3000", "")


def normalize_store_name(value: object) -> str:
    return str(value or "").strip().casefold()


def detect_encoding(path: Path) -> str:
    raw = path.read_bytes()[:4096]
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "utf-8-sig"


def is_blank_row(row: Iterable[object]) -> bool:
    return all(value is None or str(value).strip() == "" for value in row)


def has_second_row_data(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        assert load_workbook is not None
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = ws.iter_rows(min_row=2, max_row=2, values_only=True)
            row = next(rows, None)
            return bool(row and not is_blank_row(row))
        finally:
            wb.close()
    if suffix == ".csv":
        encoding = detect_encoding(path)
        with path.open("r", encoding=encoding, newline="") as file:
            reader = csv.reader(file)
            next(reader, None)
            row = next(reader, None)
            return bool(row and not is_blank_row(row))
    return True


def open_folder(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    if sys.platform.startswith("win"):
        os.startfile(str(path))  # type: ignore[attr-defined]
    else:
        subprocess.Popen(["xdg-open", str(path)])


def ensure_dependencies() -> bool:
    global Workbook, load_workbook
    if Workbook is not None and load_workbook is not None:
        return True

    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "-r", str(APP_DIR / "requirements.txt")])
        from openpyxl import Workbook as installed_workbook
        from openpyxl import load_workbook as installed_load_workbook

        Workbook = installed_workbook
        load_workbook = installed_load_workbook
        return True
    except Exception as exc:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("依赖安装失败", f"无法自动安装所需库：\n{exc}")
        root.destroy()
        return False


class GmvProcessor:
    def __init__(
        self,
        source_dir: Path,
        output_dir: Path,
        mapping_file: Path | None,
        promotion_file: Path | None,
        fallback_operator: str,
        clear_output: bool,
        log: Callable[[str], None],
    ) -> None:
        self.source_dir = source_dir
        self.output_dir = output_dir
        self.mapping_file = mapping_file
        self.promotion_file = promotion_file
        self.fallback_operator = fallback_operator.strip() or DEFAULT_OPERATOR
        self.clear_output = clear_output
        self.log = log
        self.operator_map: dict[str, str] = {}

    def load_operator_map(self) -> dict[str, str]:
        if not self.mapping_file:
            self.log("未选择运营映射表，将使用默认自定义名称。")
            self.operator_map = {}
            return {}
        if not self.mapping_file.exists():
            raise FileNotFoundError(f"运营映射表不存在：{self.mapping_file}")

        suffix = self.mapping_file.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            mapping = self._load_xlsx_mapping(self.mapping_file)
        elif suffix == ".csv":
            mapping = self._load_csv_mapping(self.mapping_file)
        else:
            raise ValueError("运营映射表仅支持 .xlsx/.xlsm/.csv")

        self.operator_map = mapping
        self.log(f"已读取运营映射：{len(mapping)} 个店铺。")
        return mapping

    def _load_xlsx_mapping(self, path: Path) -> dict[str, str]:
        assert load_workbook is not None
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                return {}
            store_index, operator_index = self._find_mapping_columns(header)
            result: dict[str, str] = {}
            for row in rows:
                store = self._cell(row, store_index)
                operator = self._cell(row, operator_index)
                if store:
                    result[normalize_store_name(store)] = operator or self.fallback_operator
            return result
        finally:
            wb.close()

    def _load_csv_mapping(self, path: Path) -> dict[str, str]:
        encoding = detect_encoding(path)
        with path.open("r", encoding=encoding, newline="") as file:
            reader = csv.reader(file)
            header = next(reader, None)
            if not header:
                return {}
            store_index, operator_index = self._find_mapping_columns(header)
            result: dict[str, str] = {}
            for row in reader:
                store = self._cell(row, store_index)
                operator = self._cell(row, operator_index)
                if store:
                    result[normalize_store_name(store)] = operator or self.fallback_operator
            return result

    def _find_mapping_columns(self, header: Iterable[object]) -> tuple[int, int]:
        labels = [normalize_header(value) for value in header]

        store_aliases = {"店铺名称", "店铺名", "店铺", "账号", "账户", "店铺账号"}
        operator_aliases = {"运营", "负责人", "人员", "自定义名称", "姓名"}

        store_index = next((i for i, label in enumerate(labels) if label in store_aliases), None)
        operator_index = next(
            (i for i, label in enumerate(labels) if label in operator_aliases), None
        )

        if store_index is None or operator_index is None:
            raise ValueError(
                "运营映射表必须包含“店铺名称/店铺名”和“运营/负责人”两列。"
            )
        return store_index, operator_index

    @staticmethod
    def _cell(row: Iterable[object], index: int) -> str:
        values = list(row)
        if index >= len(values):
            return ""
        return str(values[index] or "").strip()

    def operator_for_store(self, store: str) -> str:
        return self.operator_map.get(normalize_store_name(store), self.fallback_operator)

    def prepare_output_dir(self) -> None:
        if self.clear_output and self.output_dir.exists():
            for child in self.output_dir.iterdir():
                if child.is_dir():
                    shutil.rmtree(child)
                else:
                    child.unlink()
            self.log("已清空输出目录。")
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def scan_rule1(self) -> list[tuple[Path, FileRule, Path]]:
        if not self.source_dir.exists():
            raise FileNotFoundError(f"下载根目录不存在：{self.source_dir}")

        matches: list[tuple[Path, FileRule, Path]] = []
        store_folders = [p for p in self.source_dir.iterdir() if p.is_dir()]
        for store_folder in sorted(store_folders, key=lambda p: p.name.lower()):
            store = store_folder.name
            operator = self.operator_for_store(store)
            for source_file in sorted(store_folder.iterdir(), key=lambda p: p.name.lower()):
                if not source_file.is_file():
                    continue
                rule = self.match_rule(source_file.name)
                if not rule:
                    continue
                if not has_second_row_data(source_file):
                    self.log(f"跳过空表：{source_file.parent.name}\\{source_file.name}")
                    continue
                target = self.make_rule1_target(source_file, store, operator, rule)
                matches.append((source_file, rule, target))
        return matches

    @staticmethod
    def match_rule(file_name: str) -> FileRule | None:
        for rule in FILE_RULES:
            if rule.pattern.lower() in file_name.lower():
                return rule
        return None

    def make_rule1_target(
        self, source_file: Path, store: str, operator: str, rule: FileRule
    ) -> Path:
        suffix = source_file.suffix
        file_name = rule.template.format(
            store=safe_filename(store),
            operator=safe_filename(operator),
            stem=safe_filename(source_file.stem),
            suffix=suffix,
        )
        return self.output_dir / safe_filename(operator) / rule.output_folder / file_name

    def process_rule1(self) -> int:
        matches = self.scan_rule1()
        if not matches:
            self.log("规则1：没有找到可归档的下载文件。")
            return 0

        count = 0
        for source_file, rule, target in matches:
            target.parent.mkdir(parents=True, exist_ok=True)
            final_path = unique_path(target)
            shutil.copy2(source_file, final_path)
            count += 1
            self.log(f"规则1：{source_file.parent.name}\\{source_file.name} -> {rule.output_folder}\\{final_path.name}")
        self.log(f"规则1完成：共归档 {count} 个文件。")
        return count

    def process_promotion_split(self) -> int:
        if not self.promotion_file:
            self.log("未选择推广明细文件，跳过规则2。")
            return 0
        if not self.promotion_file.exists():
            raise FileNotFoundError(f"推广明细文件不存在：{self.promotion_file}")
        if self.promotion_file.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValueError("推广明细拆分仅支持 .xlsx/.xlsm 文件。")
        if not self.operator_map:
            self.log("规则2：没有运营匹配表，跳过推广明细拆分。")
            return 0
        if not has_second_row_data(self.promotion_file):
            self.log("规则2：推广订单表只有表头，跳过拆分。")
            return 0

        assert load_workbook is not None and Workbook is not None
        wb = load_workbook(self.promotion_file, read_only=True, data_only=True)
        groups: dict[str, list[tuple[object, ...]]] = {}
        header: tuple[object, ...] | None = None
        try:
            ws = wb.active
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = tuple(row)
                if row_index == 1:
                    header = values
                    continue
                if is_blank_row(values):
                    continue
                store = str(values[1] or "").strip() if len(values) > 1 else ""
                if not store:
                    continue
                if normalize_store_name(store) not in self.operator_map:
                    continue
                groups.setdefault(store, []).append(values)
        finally:
            wb.close()

        if not header:
            raise ValueError("推广明细文件没有表头。")
        if not groups:
            self.log("规则2：推广明细没有可拆分的数据。")
            return 0

        for store, rows in sorted(groups.items(), key=lambda item: item[0].lower()):
            operator = self.operator_for_store(store)
            target_dir = self.output_dir / safe_filename(operator) / "推广明细"
            target_dir.mkdir(parents=True, exist_ok=True)
            file_name = safe_filename(f"{store}-推广明细-{operator}.xlsx")
            target = unique_path(target_dir / file_name)

            new_wb = Workbook()
            ws = new_wb.active
            ws.title = "Sheet1"
            ws.append(header)
            for row in rows:
                ws.append(row)
            self._autosize_columns(ws)
            new_wb.save(target)
            self.log(f"规则2：{store} -> 推广明细\\{target.name}（{len(rows)} 行）")

        self.log(f"规则2完成：共拆分 {len(groups)} 个店铺。")
        return len(groups)

    @staticmethod
    def _autosize_columns(ws) -> None:
        for column_cells in ws.columns:
            first_cell = column_cells[0]
            max_length = 8
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 40))
            ws.column_dimensions[first_cell.column_letter].width = max_length + 2

    def run(self) -> tuple[int, int]:
        if load_workbook is None or Workbook is None:
            raise RuntimeError(
                "缺少 openpyxl 依赖。请先运行 install_requirements.bat 安装依赖。"
            )
        self.load_operator_map()
        self.prepare_output_dir()
        rule1_count = self.process_rule1()
        promotion_count = self.process_promotion_split()
        return rule1_count, promotion_count


class App(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("自动 GMV 文件整理 1.0")
        self.geometry("1080x720")
        self.minsize(980, 640)

        self.log_queue: queue.Queue[tuple[str, str]] = queue.Queue()
        self.worker: threading.Thread | None = None

        self.source_var = tk.StringVar(value=str(DEFAULT_SOURCE_DIR))
        self.output_var = tk.StringVar(value=str(DEFAULT_OUTPUT_DIR))
        self.mapping_var = tk.StringVar(value="")
        self.promotion_var = tk.StringVar(value="")
        self.operator_var = tk.StringVar(value=DEFAULT_OPERATOR)
        self.clear_output_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="准备就绪")

        self._load_config()
        self._setup_style()
        self._build_ui()
        self._poll_log_queue()

    def _setup_style(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        bg = "#f6f8fb"
        panel = "#ffffff"
        accent = "#2563eb"
        text = "#1f2937"

        self.configure(bg=bg)
        style.configure(".", font=("Microsoft YaHei UI", 10), background=bg, foreground=text)
        style.configure("TFrame", background=bg)
        style.configure("Panel.TFrame", background=panel)
        style.configure("TLabel", background=bg, foreground=text)
        style.configure("Panel.TLabel", background=panel, foreground=text)
        style.configure("Title.TLabel", font=("Microsoft YaHei UI", 18, "bold"), background=bg)
        style.configure("Hint.TLabel", foreground="#64748b", background=bg)
        style.configure("PanelHint.TLabel", foreground="#64748b", background=panel)
        style.configure("TButton", padding=(12, 7))
        style.configure("Accent.TButton", background=accent, foreground="#ffffff")
        style.map("Accent.TButton", background=[("active", "#1d4ed8")])
        style.configure("TEntry", padding=(8, 5))
        style.configure("Treeview", rowheight=28, font=("Microsoft YaHei UI", 9))
        style.configure("Treeview.Heading", font=("Microsoft YaHei UI", 10, "bold"))

    def _build_ui(self) -> None:
        root = ttk.Frame(self, padding=18)
        root.pack(fill=tk.BOTH, expand=True)

        header = ttk.Frame(root)
        header.pack(fill=tk.X)
        ttk.Label(header, text="自动 GMV 文件整理", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Label(
            header,
            text="按店铺文件夹归档下载表，并按 B 列拆分推广明细",
            style="Hint.TLabel",
        ).pack(side=tk.LEFT, padx=(18, 0), pady=(8, 0))

        main = ttk.PanedWindow(root, orient=tk.VERTICAL)
        main.pack(fill=tk.BOTH, expand=True, pady=(16, 0))

        top_panel = ttk.Frame(main, style="Panel.TFrame", padding=16)
        bottom_panel = ttk.Frame(main, style="Panel.TFrame", padding=12)
        main.add(top_panel, weight=0)
        main.add(bottom_panel, weight=1)

        self._build_form(top_panel)
        self._build_preview(bottom_panel)

        status_bar = ttk.Frame(root)
        status_bar.pack(fill=tk.X, pady=(10, 0))
        ttk.Label(status_bar, textvariable=self.status_var, style="Hint.TLabel").pack(side=tk.LEFT)
        self.progress = ttk.Progressbar(status_bar, mode="indeterminate", length=180)
        self.progress.pack(side=tk.RIGHT)

    def _build_form(self, parent: ttk.Frame) -> None:
        parent.columnconfigure(1, weight=1)

        self._path_row(
            parent,
            row=0,
            label="源数据文件夹",
            variable=self.source_var,
            command=self._choose_source_dir,
            hint="选择包含各店铺子文件夹的目录，例如 ziniao browser。",
        )
        self._path_row(
            parent,
            row=1,
            label="目标输出文件夹",
            variable=self.output_var,
            command=lambda: self._choose_dir(self.output_var),
            hint="整理后的订单明细、广告明细、广告消耗明细、退款明细、推广明细会放在这里。",
        )
        self._path_row(
            parent,
            row=2,
            label="选择店铺-运营 匹配表(可导出模版)",
            variable=self.mapping_var,
            command=lambda: self._choose_file(
                self.mapping_var, [("表格文件", "*.xlsx *.xlsm *.csv"), ("所有文件", "*.*")]
            ),
            hint="可选。表头需包含“店铺名称/店铺名”和“运营/负责人”。",
            extra_text="导出模版",
            extra_command=self.export_mapping_template,
        )
        self._path_row(
            parent,
            row=3,
            label="选择推广订单表",
            variable=self.promotion_var,
            command=lambda: self._choose_file(
                self.promotion_var, [("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")]
            ),
            hint="可选。会按 B 列店铺名称拆分，保留第 1 行表头。",
        )

        button_row = ttk.Frame(parent, style="Panel.TFrame")
        button_row.grid(row=4, column=0, columnspan=3, sticky=tk.EW, pady=(16, 0))
        ttk.Button(button_row, text="扫描预览", command=self.scan_preview).pack(side=tk.LEFT)
        ttk.Button(button_row, text="开始整理", style="Accent.TButton", command=self.start).pack(
            side=tk.LEFT, padx=(10, 0)
        )
        ttk.Button(button_row, text="打开输出目录", command=self.open_output).pack(
            side=tk.LEFT, padx=(10, 0)
        )

    def _path_row(
        self,
        parent: ttk.Frame,
        row: int,
        label: str,
        variable: tk.StringVar,
        command: Callable[[], None],
        hint: str,
        extra_text: str | None = None,
        extra_command: Callable[[], None] | None = None,
    ) -> None:
        ttk.Label(parent, text=label, style="Panel.TLabel").grid(
            row=row, column=0, sticky=tk.W, pady=(0, 10)
        )
        frame = ttk.Frame(parent, style="Panel.TFrame")
        frame.grid(row=row, column=1, sticky=tk.EW, pady=(0, 10))
        frame.columnconfigure(0, weight=1)
        ttk.Entry(frame, textvariable=variable).grid(row=0, column=0, sticky=tk.EW)
        ttk.Button(frame, text="选择", command=command).grid(row=0, column=1, padx=(8, 0))
        if extra_text and extra_command:
            ttk.Button(frame, text=extra_text, command=extra_command).grid(
                row=0, column=2, padx=(8, 0)
            )
        ttk.Label(parent, text=hint, style="PanelHint.TLabel").grid(
            row=row, column=2, sticky=tk.W, padx=(12, 0), pady=(0, 10)
        )

    def _build_preview(self, parent: ttk.Frame) -> None:
        parent.rowconfigure(0, weight=1)
        parent.columnconfigure(0, weight=1)

        preview_frame = ttk.Frame(parent, padding=8)
        preview_frame.grid(row=0, column=0, sticky=tk.NSEW)

        preview_frame.rowconfigure(0, weight=1)
        preview_frame.columnconfigure(0, weight=1)

        columns = ("store", "rule", "source", "target")
        self.preview_tree = ttk.Treeview(
            preview_frame, columns=columns, show="headings", selectmode="browse"
        )
        self.preview_tree.heading("store", text="店铺")
        self.preview_tree.heading("rule", text="归档目录")
        self.preview_tree.heading("source", text="来源文件")
        self.preview_tree.heading("target", text="目标文件")
        self.preview_tree.column("store", width=150, anchor=tk.W)
        self.preview_tree.column("rule", width=120, anchor=tk.W)
        self.preview_tree.column("source", width=320, anchor=tk.W)
        self.preview_tree.column("target", width=360, anchor=tk.W)
        self.preview_tree.grid(row=0, column=0, sticky=tk.NSEW)

        y_scroll = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL, command=self.preview_tree.yview)
        y_scroll.grid(row=0, column=1, sticky=tk.NS)
        self.preview_tree.configure(yscrollcommand=y_scroll.set)

    def _choose_dir(self, variable: tk.StringVar) -> None:
        initial = variable.get().strip() or str(APP_DIR)
        path = filedialog.askdirectory(initialdir=initial)
        if path:
            variable.set(path)

    def _choose_source_dir(self) -> None:
        self._choose_dir(self.source_var)
        self.save_config()

    def _choose_file(self, variable: tk.StringVar, filetypes) -> None:
        initial_value = variable.get().strip()
        initial = str(Path(initial_value).parent) if initial_value else str(APP_DIR)
        path = filedialog.askopenfilename(initialdir=initial, filetypes=filetypes)
        if path:
            variable.set(path)

    def _processor(self, log: Callable[[str], None] | None = None) -> GmvProcessor:
        mapping = self.mapping_var.get().strip()
        promotion = self.promotion_var.get().strip()
        return GmvProcessor(
            source_dir=Path(self.source_var.get().strip()),
            output_dir=Path(self.output_var.get().strip()),
            mapping_file=Path(mapping) if mapping else None,
            promotion_file=Path(promotion) if promotion else None,
            fallback_operator=self.operator_var.get(),
            clear_output=self.clear_output_var.get(),
            log=log or self.log,
        )

    def scan_preview(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("正在运行", "任务执行中，请稍后再扫描。")
            return

        self.preview_tree.delete(*self.preview_tree.get_children())
        try:
            processor = self._processor()
            processor.load_operator_map()
            matches = processor.scan_rule1()
        except Exception as exc:
            messagebox.showerror("扫描失败", str(exc))
            self.status_var.set("扫描失败")
            return

        for source_file, rule, target in matches:
            self.preview_tree.insert(
                "",
                tk.END,
                values=(
                    source_file.parent.name,
                    f"{target.parent.parent.name}\\{rule.output_folder}",
                    source_file.name,
                    str(target),
                ),
            )
        self.status_var.set(f"扫描完成：规则1可归档 {len(matches)} 个文件")
        self.log(f"扫描完成：发现 {len(matches)} 个规则1文件。")

    def start(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("正在运行", "整理任务已经在运行。")
            return

        if not self.source_var.get().strip():
            messagebox.showwarning("缺少路径", "请先选择源数据文件夹。")
            return
        if not self.output_var.get().strip():
            messagebox.showwarning("缺少路径", "请先选择目标输出文件夹。")
            return

        self.save_config(show_message=False)
        self.status_var.set("正在整理...")
        self.progress.start(10)

        def work() -> None:
            try:
                processor = self._processor(log=lambda message: None)
                rule1_count, promotion_count = processor.run()
                self.log_queue.put(
                    (
                        "done",
                        f"整理完成：规则1归档 {rule1_count} 个文件，规则2拆分 {promotion_count} 个店铺。",
                    )
                )
            except Exception as exc:
                self.log_queue.put(("error", str(exc)))

        self.worker = threading.Thread(target=work, daemon=True)
        self.worker.start()

    def _poll_log_queue(self) -> None:
        try:
            while True:
                level, message = self.log_queue.get_nowait()
                if level == "log":
                    pass
                elif level == "done":
                    self.progress.stop()
                    self.worker = None
                    self.status_var.set(message)
                    messagebox.showinfo("完成", message)
                    self.scan_preview()
                elif level == "error":
                    self.progress.stop()
                    self.worker = None
                    self.status_var.set("整理失败")
                    messagebox.showerror("整理失败", message)
        except queue.Empty:
            pass
        self.after(120, self._poll_log_queue)

    def log(self, message: str) -> None:
        self.status_var.set(message.rstrip())

    def export_mapping_template(self) -> None:
        if Workbook is None:
            messagebox.showerror("缺少依赖", "正在自动安装依赖，请重新启动程序后再导出模版。")
            return
        default_path = APP_DIR / "店铺-运营匹配表模版.xlsx"
        path = filedialog.asksaveasfilename(
            initialdir=str(APP_DIR),
            initialfile=default_path.name,
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "店铺运营匹配"
        ws.append(["店铺名称", "运营"])
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 18
        wb.save(path)
        self.mapping_var.set(path)
        self.save_config(show_message=False)
        messagebox.showinfo("已导出", f"匹配表模版已导出：\n{path}")

    def open_output(self) -> None:
        try:
            open_folder(Path(self.output_var.get().strip()))
        except Exception as exc:
            messagebox.showerror("打开失败", str(exc))

    def _load_config(self) -> None:
        if not CONFIG_PATH.exists():
            return
        try:
            data = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except Exception:
            return
        self.source_var.set(data.get("source_dir", self.source_var.get()))

    def save_config(self, show_message: bool = False) -> None:
        data = {
            "source_dir": self.source_var.get().strip(),
        }
        CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        if show_message:
            messagebox.showinfo("已保存", f"配置已保存到：{CONFIG_PATH}")


def main() -> None:
    if not ensure_dependencies():
        return

    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
